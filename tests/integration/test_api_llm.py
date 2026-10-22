from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from jobpilot.agents.profile.schema import CandidateProfile
from jobpilot.api.app import create_app
from jobpilot.core.config import AppConfig, ModelRole, Settings
from jobpilot.core.context import AppContext
from jobpilot.core.crypto import Vault
from jobpilot.core.migrate import upgrade_to_head
from jobpilot.llm.base import LLMError
from tests.conftest import FakeLLM

AUTH = {"Authorization": "Bearer t"}


@pytest.fixture
def ctx(tmp_path: Path) -> Iterator[AppContext]:
    settings = Settings(
        _env_file=None,
        data_dir=tmp_path,
        jobpilot_api_token="t",
        jobpilot_master_key=Vault.generate_key(),
        database_url=f"sqlite:///{tmp_path / 'api.db'}",
    )
    context = AppContext(settings=settings, config=AppConfig.load(Path("config.yaml")))
    upgrade_to_head(context.engine)
    yield context
    context.engine.dispose()


@pytest.fixture
def client(ctx: AppContext) -> TestClient:
    return TestClient(create_app(ctx, web_dist=None))


def test_presets_include_open_source_and_paid(client: TestClient) -> None:
    presets = {p["id"]: p for p in client.get("/api/v1/llm/presets", headers=AUTH).json()["data"]}
    assert presets["ollama"]["is_local"] and presets["ollama"]["structured_mode"] == "json_object"
    assert {"anthropic", "openai", "gemini", "groq", "openrouter", "litellm"} <= set(presets)


def test_add_provider_route_and_never_leak_key(client: TestClient) -> None:
    saved = client.put(
        "/api/v1/llm/providers/groq",
        headers=AUTH,
        json={
            "preset": "groq",
            "type": "openai_compatible",
            "base_url": "https://api.groq.com/openai/v1",
            "structured_mode": "json_object",
            "api_key": "gsk-very-secret",
        },
    )
    assert saved.status_code == 200 and saved.json()["data"]["has_api_key"] is True
    listing = client.get("/api/v1/llm/providers", headers=AUTH)
    assert "gsk-very-secret" not in listing.text

    route = client.put(
        "/api/v1/llm/routes/writing",
        headers=AUTH,
        json={
            "provider": "groq",
            "model": "llama-3.3-70b-versatile",
            "fallbacks": [{"provider": "ollama", "model": "gemma3:12b"}],
        },
    )
    assert route.json()["data"]["source"] == "user"
    roles = {r["role"]: r for r in client.get("/api/v1/llm/routes", headers=AUTH).json()["data"]}
    assert roles["writing"]["provider"] == "groq" and roles["fast"]["provider"] == "ollama"
    reset = client.delete("/api/v1/llm/routes/writing", headers=AUTH).json()["data"]
    assert reset["provider"] == "ollama"


def test_settings_errors_are_400_and_rolled_back(client: TestClient) -> None:
    bad = client.put(
        "/api/v1/llm/routes/fast", headers=AUTH, json={"provider": "ghost", "model": "x"}
    )
    assert bad.status_code == 400 and "ghost" in bad.json()["error"]["message"]
    assert client.get("/api/v1/llm/settings", headers=AUTH).json()["data"]["revision"] == 0


def test_bad_provider_name_is_rejected(client: TestClient) -> None:
    response = client.put(
        "/api/v1/llm/providers/Bad%20Name",
        headers=AUTH,
        json={"type": "openai_compatible", "base_url": "http://x"},
    )
    assert response.status_code == 422


def test_privacy_mode_change_reaches_tracer_check_immediately(
    client: TestClient, ctx: AppContext
) -> None:
    assert ctx.privacy.local_only() is False  # primes the cache
    client.put("/api/v1/llm/settings/privacy", headers=AUTH, json={"privacy_mode": "local_only"})
    assert ctx.privacy.local_only() is True  # no waiting for the cache TTL


def test_privacy_mode_round_trip(client: TestClient) -> None:
    out = client.put(
        "/api/v1/llm/settings/privacy", headers=AUTH, json={"privacy_mode": "local_only"}
    ).json()["data"]
    assert out["privacy_mode"] == "local_only" and out["vault_ready"] is True
    assert client.get("/api/v1/health").json()["data"]["llm_privacy_mode"] == "local_only"


def test_models_and_connection_test(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "jobpilot.llm.gateway.LazyProvider.list_models", lambda self: ["qwen3:8b", "gemma3:12b"]
    )
    models = client.get("/api/v1/llm/providers/ollama/models", headers=AUTH).json()
    assert models["data"]["models"] == ["qwen3:8b", "gemma3:12b"]
    test = client.post("/api/v1/llm/providers/ollama/test", headers=AUTH).json()["data"]
    assert test["ok"] and test["models_found"] == 2


def test_connection_test_reports_failure(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    def boom(self):  # noqa: ANN001, ANN202
        raise LLMError("could not reach ollama")

    monkeypatch.setattr("jobpilot.llm.gateway.LazyProvider.list_models", boom)
    test = client.post("/api/v1/llm/providers/ollama/test", headers=AUTH).json()["data"]
    assert test == test | {"ok": False, "error": "could not reach ollama"}
    assert client.get("/api/v1/llm/providers/ollama/models", headers=AUTH).status_code == 502


def test_try_role_uses_route(
    client: TestClient, ctx: AppContext, monkeypatch: pytest.MonkeyPatch
) -> None:
    class Pong(CandidateProfile):  # any BaseModel works for the fake
        pass

    calls: list[ModelRole] = []

    def fake(**kwargs):  # noqa: ANN003, ANN202
        calls.append(kwargs["role"])
        from jobpilot.llm.base import StructuredResult, Usage

        return StructuredResult(
            kwargs["schema"](reply="hello from qwen"), Usage(5, 3), "ollama", "qwen3:8b"
        )

    monkeypatch.setattr(ctx.gateway, "structured", fake)
    out = client.post("/api/v1/llm/routes/extraction/try", headers=AUTH).json()["data"]
    assert out["ok"] and out["reply"] == "hello from qwen" and calls == [ModelRole.EXTRACTION]


def test_profile_import_upload(
    client: TestClient, ctx: AppContext, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        AppContext, "gateway", FakeLLM(CandidateProfile(full_name="Asha Rao", email="a@x.com"))
    )
    resume = tmp_path / "cv.docx"
    document = docx.Document()
    document.add_paragraph("Asha Rao, Backend Engineer")
    document.save(str(resume))
    with resume.open("rb") as fh:
        response = client.post(
            "/api/v1/profile/import",
            headers=AUTH,
            files={"file": ("cv.docx", fh)},
            data={"label": "main"},
        )
    body = response.json()
    assert response.status_code == 200, body
    assert body["data"]["version"] == 1 and body["data"]["resume_label"] == "main"
    assert any(m["field"] == "phone" for m in body["data"]["missing"])
    assert {"email", "full_name"} <= set(body["data"]["updated"]) and body["data"]["kept"] == []


def test_profile_import_reports_kept_fields_and_runs_one_at_a_time(
    client: TestClient, ctx: AppContext, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from jobpilot.agents.profile.repository import ProfileRepository
    from jobpilot.api.routes import profile as profile_routes
    from jobpilot.core.db import session_scope

    with session_scope(ctx.sessions) as s:
        ProfileRepository(s).save(CandidateProfile(full_name="Asha Rao"), source="qa")
    monkeypatch.setattr(
        AppContext, "gateway", FakeLLM(CandidateProfile(full_name="Juned Alam", email="j@x.com"))
    )
    resume = tmp_path / "cv.docx"
    document = docx.Document()
    document.add_paragraph("Juned Alam, Engineer")
    document.save(str(resume))

    assert profile_routes._IMPORT_LOCK.acquire(blocking=False)  # an import is running
    try:
        with resume.open("rb") as fh:
            busy = client.post(
                "/api/v1/profile/import", headers=AUTH, files={"file": ("cv.docx", fh)}
            )
        assert busy.status_code == 409 and "already being read" in busy.json()["error"]["message"]
    finally:
        profile_routes._IMPORT_LOCK.release()

    with resume.open("rb") as fh:
        data = client.post(
            "/api/v1/profile/import", headers=AUTH, files={"file": ("cv.docx", fh)}
        ).json()["data"]
    # The confirmed name wins and is reported, so the user can see why it didn't change.
    assert data["kept"] == ["full_name"] and data["updated"] == ["email"]


def test_profile_import_rejects_wrong_type(client: TestClient) -> None:
    response = client.post(
        "/api/v1/profile/import", headers=AUTH, files={"file": ("cv.txt", b"hello")}
    )
    assert response.status_code == 400


def test_ui_is_served_when_built(ctx: AppContext, tmp_path: Path) -> None:
    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<!doctype html><title>JobPilot</title>")
    (dist / "assets" / "app.js").write_text("console.log(1)")
    client = TestClient(create_app(ctx, web_dist=dist))
    assert "JobPilot" in client.get("/models").text  # client-side route → index.html
    assert client.get("/assets/app.js").text == "console.log(1)"
    assert client.get("/api/v1/nope").status_code == 404
    assert "JobPilot" in client.get("/../../etc/passwd").text  # traversal falls back to index


def test_profile_qa_parses_free_text(client: TestClient) -> None:
    response = client.post(
        "/api/v1/profile/qa",
        headers=AUTH,
        json={
            "answers": {
                "expected_ctc": "30 LPA",
                "notice_period_days": "60 days",
                "preferred_locations": "Pune, Delhi",
                "full_name": "Asha Rao",
                "phone": "  ",
            }
        },
    )
    data = response.json()["data"]
    assert data["expected_ctc"]["amount"] == 3_000_000 and data["notice_period_days"] == 60
    assert data["preferred_locations"] == ["Pune", "Delhi"] and data["phone"] == ""


def test_profile_qa_reports_bad_answers(client: TestClient) -> None:
    response = client.post(
        "/api/v1/profile/qa",
        headers=AUTH,
        json={"answers": {"work_modes": "sometimes", "notice_period_days": "soon"}},
    )
    message = response.json()["error"]["message"]
    assert (
        response.status_code == 400 and "work_modes" in message and "notice_period_days" in message
    )
