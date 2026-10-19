"""Resume text extraction for PDF and DOCX files."""

from __future__ import annotations

from pathlib import Path

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx"})
MAX_RESUME_BYTES = 10 * 1024 * 1024


class ResumeParseError(ValueError):
    pass


def extract_text(path: Path) -> str:
    if not path.is_file():
        raise ResumeParseError(f"resume not found: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ResumeParseError(f"unsupported resume type '{suffix}', use PDF or DOCX")
    if path.stat().st_size > MAX_RESUME_BYTES:
        raise ResumeParseError("resume is larger than 10 MB")

    try:
        text = _pdf_text(path) if suffix == ".pdf" else _docx_text(path)
    except (PdfReadError, ValueError, KeyError) as exc:
        raise ResumeParseError(f"could not read {path.name}: {exc}") from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise ResumeParseError(f"no text found in {path.name} (scanned image PDFs need OCR)")
    return text


def _pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
